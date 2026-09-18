import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import database

STATUS_TR = {
    "PAID": "Tamamı Ödendi",
    "PARTIAL": "Alacak Var (Kısmi)",
    "PENDING": "Ödeme Bekliyor"
}

def generate_weekly_excel_report(start_date_str, end_date_str, firm_name, filepath):
    records = database.get_weekly_records(start_date_str, end_date_str, firm_name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Haftalık Yövmiye Raporu"

    # Renk Tanımları
    navy_fill = PatternFill(start_color="0D2C3A", end_color="0D2C3A", fill_type="solid")
    orange_fill = PatternFill(start_color="D9531E", end_color="D9531E", fill_type="solid")
    light_gray_fill = PatternFill(start_color="F4F6F8", end_color="F4F6F8", fill_type="solid")
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")

    font_title = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10)
    font_total = Font(name="Segoe UI", size=11, bold=True, color="0D2C3A")

    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    thick_bottom = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='medium', color='0D2C3A'),
        bottom=Side(style='double', color='0D2C3A')
    )

    # Başlık Alanı (Satır 1 ve 2)
    ws.merge_cells("A1:H1")
    ws["A1"] = "MD FUAR HİZMETLERİ - YÖVMİYE VE HAKEDİŞ RAPORU"
    ws["A1"].font = font_title
    ws["A1"].fill = navy_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    firm_title_part = f" | Firma Filteresi: {firm_name}" if firm_name else ""
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Rapor Tarihi: {start_date_str} - {end_date_str}{firm_title_part}"
    ws["A2"].font = font_subtitle
    ws["A2"].fill = navy_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Boş Satır
    ws.row_dimensions[3].height = 10

    # Tablo Başlıkları (Satır 4)
    headers = [
        "Tarih", "Firma / Fuar Adı", "Yövmiye Sayısı", 
        "Alınan Ödeme (₺)", "İşçilere Ödenen (₺)", "Net Kazanç (₺)", 
        "Ödeme Durumu", "Notlar & Detaylar"
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header)
        cell.font = font_header
        cell.fill = orange_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 28

    # Verileri Yazma
    start_row = 5
    total_yovmiye = 0.0
    total_received = 0.0
    total_expense = 0.0
    total_net = 0.0

    for idx, r in enumerate(records):
        row_num = start_row + idx
        
        yov = float(r["yovmiye"] or 0)
        recv = float(r["amount_received"] or 0)
        exp = float(r["worker_expense"] or 0)
        net = float(r["net_profit"] or 0)
        status = STATUS_TR.get(r.get("payment_status", "PAID"), "Tamamı Ödendi")

        total_yovmiye += yov
        total_received += recv
        total_expense += exp
        total_net += net

        ws.cell(row=row_num, column=1, value=r["record_date"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=2, value=r["firm_name"]).alignment = Alignment(horizontal="left")
        
        c_yov = ws.cell(row=row_num, column=3, value=yov)
        c_yov.number_format = "#,##0.0"
        c_yov.alignment = Alignment(horizontal="center")

        c_recv = ws.cell(row=row_num, column=4, value=recv)
        c_recv.number_format = "#,##0.00 TL"
        c_recv.alignment = Alignment(horizontal="right")

        c_exp = ws.cell(row=row_num, column=5, value=exp)
        c_exp.number_format = "#,##0.00 TL"
        c_exp.alignment = Alignment(horizontal="right")

        c_net = ws.cell(row=row_num, column=6, value=net)
        c_net.number_format = "#,##0.00 TL"
        c_net.alignment = Alignment(horizontal="right")

        ws.cell(row=row_num, column=7, value=status).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=8, value=r["notes"]).alignment = Alignment(horizontal="left")

        for c in range(1, 9):
            cell = ws.cell(row=row_num, column=c)
            cell.font = font_data
            cell.border = thin_border
            if idx % 2 == 1:
                cell.fill = light_gray_fill

        ws.row_dimensions[row_num].height = 22

    # Toplam Satırı
    tot_row = start_row + len(records)
    ws.cell(row=tot_row, column=1, value="GENEL TOPLAM").alignment = Alignment(horizontal="center")
    ws.cell(row=tot_row, column=2, value=f"{len(records)} Kayıt").alignment = Alignment(horizontal="left")
    
    c_tyov = ws.cell(row=tot_row, column=3, value=total_yovmiye)
    c_tyov.number_format = "#,##0.0"
    c_tyov.alignment = Alignment(horizontal="center")

    c_trecv = ws.cell(row=tot_row, column=4, value=total_received)
    c_trecv.number_format = "#,##0.00 TL"
    c_trecv.alignment = Alignment(horizontal="right")

    c_texp = ws.cell(row=tot_row, column=5, value=total_expense)
    c_texp.number_format = "#,##0.00 TL"
    c_texp.alignment = Alignment(horizontal="right")

    c_tnet = ws.cell(row=tot_row, column=6, value=total_net)
    c_tnet.number_format = "#,##0.00 TL"
    c_tnet.alignment = Alignment(horizontal="right")

    ws.cell(row=tot_row, column=7, value="").alignment = Alignment(horizontal="center")
    ws.cell(row=tot_row, column=8, value="").alignment = Alignment(horizontal="left")

    for c in range(1, 9):
        cell = ws.cell(row=tot_row, column=c)
        cell.font = font_total
        cell.fill = total_fill
        cell.border = thick_bottom

    ws.row_dimensions[tot_row].height = 26

    # Sütun Genişliklerini Otomatik Ayarla
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(filepath)
    return len(records)
